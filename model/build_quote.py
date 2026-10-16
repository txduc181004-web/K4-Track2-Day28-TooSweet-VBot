"""Build the quote workbook (G??_baogia.xlsx) from model/vbot_model.py.

usage: .venv/bin/python model/build_quote.py [GROUP_CODE]   (default G00)
Sheets: BaoGia (5 columns, last row TOTAL), BangChung (evidence E01..), GiaDinh (sizing for reviewers).
Costs are written as numbers (not formulas) so the checker reads them without recalculation.
"""
import os, sys
from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import vbot_model as M

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
GROUP = sys.argv[1] if len(sys.argv) > 1 else 'G00'
CAPTURE_DATE = '2026-10-16'
COMPONENTS = ['compute', 'llm', 'embedding', 'gateway', 'load_balancer', 'authn', 'authz', 'database', 'vector_db',
              'object_storage', 'cache_queue', 'sandbox', 'connectors', 'secrets_kms', 'backup_dr', 'monitoring',
              'governance', 'network', 'cicd', 'ops_labor', 'licenses', 'contingency', 'other']
ORDER = {c: i for i, c in enumerate(COMPONENTS)}

rows, S = M.build()
rows = sorted(rows, key=lambda r: ORDER[r['component']])  # stable: keeps model order inside a component
total = round(sum(r['chi_phi'] for r in rows), 2)

wb = Workbook()
ws = wb.active
ws.title = 'BaoGia'
head = ['Component', 'Ý nghĩa', 'Ghi chú', 'Chi phí/tháng (USD)', 'Bằng chứng']
ws.append(head)
for r in rows:
    ws.append([r['component'], r['y_nghia'], r['ghi_chu'], r['chi_phi'], r['bang_chung']])
ws.append(['TOTAL', 'Tổng chi phí vận hành 1 tháng (prod + staging), USD chưa VAT', f'{len(rows)} dòng chi phí; '
           f'{total / S["req_month"] * 1000:.2f} USD / 1.000 request; {total / M.A["dau"]:.2f} USD / DAU', total, ''])

hdr_fill = PatternFill('solid', fgColor='1F3A5F')
for c in ws[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = hdr_fill
    c.alignment = Alignment(vertical='center', wrap_text=True)
for row in ws.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=True)
    row[3].number_format = '#,##0.00'
last = ws.max_row
for c in ws[last]:
    c.font = Font(bold=True)
    c.fill = PatternFill('solid', fgColor='E8EEF6')
for col, w in zip('ABCDE', (16, 46, 90, 16, 18)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'
dv = DataValidation(type='list', formula1='"' + ','.join(COMPONENTS) + ',TOTAL"', allow_blank=False)
ws.add_data_validation(dv)
dv.add(f'A2:A{last}')

# ---- BangChung ----
ev = wb.create_sheet('BangChung')
ev.append(['evidence_id', 'Mô tả', 'Link', 'Ngày chụp', 'Ảnh chụp màn hình'])
shots = os.path.join(ROOT, 'evidence', 'screenshots')
for i, (eid, (desc, url)) in enumerate(M.EVIDENCE.items(), start=2):
    ev.append([eid, desc, url, CAPTURE_DATE, ''])
    png = os.path.join(shots, f'{eid}.png')
    if os.path.exists(png):
        img = XLImage(png)
        scale = 520 / img.width
        img.width, img.height = 520, int(img.height * scale)
        ev.add_image(img, f'E{i}')
        ev.row_dimensions[i].height = img.height * 0.75 + 6
for c in ev[1]:
    c.font = Font(bold=True, color='FFFFFF')
    c.fill = hdr_fill
for row in ev.iter_rows(min_row=2):
    for c in row:
        c.alignment = Alignment(vertical='top', wrap_text=True)
for col, w in zip('ABCDE', (12, 60, 48, 12, 76)):
    ev.column_dimensions[col].width = w

# ---- GiaDinh (sizing, for reviewers; not read by the checker) ----
gd = wb.create_sheet('GiaDinh')
gd.append(['Đại lượng', 'Giá trị', 'Công thức / căn cứ'])
A = M.A
facts = [
    ('DAU', A['dau'], 'Đề cho'),
    ('Request / user / ngày', A['req_per_user_day'], 'Giả định nhóm (khớp REQ_PER_USER trên trang đề)'),
    ('Request / ngày', S['req_day'], 'DAU x 25'),
    ('Request / tháng', S['req_month'], 'x 22 ngày làm việc'),
    ('Request/giây giờ cao điểm', round(S['peak_rps'], 2), '18% request/ngày trong 1 giờ / 3600'),
    ('Request/giây đỉnh theo phút', round(S['burst_rps'], 2), 'x hệ số burst 2'),
    ('Phiên streaming đồng thời (đỉnh)', round(S['concurrent_streams']), 'RPS đỉnh x 10 giây/phiên'),
    ('Token vào / giây (cao điểm)', round(S['peak_in_tps']), 'RPS cao điểm x token vào TB/request'),
    ('Token ra / giây (cao điểm)', round(S['peak_out_tps']), 'RPS cao điểm x token ra TB/request'),
    ('Token vào / tháng (tỷ)', round(S['tin_month'] / 1e9, 2), 'Tổng theo 5 cấp, gồm RAG và các bước agent'),
    ('Token ra / tháng (tỷ)', round(S['tout_month'] / 1e9, 2), ''),
    ('Lượt sandbox / tháng', round(S['sandbox_runs']), 'cấp 2 x 0.5 + cấp 3 x 2.5 + cấp 4 x 0.2 + cấp 5 x 1'),
    ('Lệnh gọi connector / tháng', round(S['conn_calls']), 'theo số lệnh gọi trung bình mỗi cấp'),
    ('Vector KB', S['kb_chunks'], '300.000 tài liệu x 8 chunk'),
    ('Workspace (GiB)', A['dau'] * M.STORAGE['workspace_gib_per_user'], 'TB 1 GB/user, quota 10 GB'),
    ('Lịch sử chat 12 tháng (GiB)', round(S['chat_gib'], 1), '8 KB/request x 12 tháng'),
]
for f in facts:
    gd.append(list(f))
gd.append([])
gd.append(['Cấp', 'Tỉ lệ request', 'Request/tháng', 'Chi phí LLM (USD)', 'Model'])
for t in S['tiers']:
    tier = next(x for x in M.TIERS if x['id'] == t['id'])
    gd.append([f"{t['id']} {t['name']}", tier['share'], round(t['req']), round(t['cost'], 2),
               ', '.join(sorted({c['model'] for c in tier['calls']}))])
for c in gd[1]:
    c.font = Font(bold=True)
for col, w in zip('ABCDE', (34, 16, 60, 18, 24)):
    gd.column_dimensions[col].width = w

out = os.path.join(ROOT, 'deliverables', f'{GROUP}_baogia.xlsx')
os.makedirs(os.path.dirname(out), exist_ok=True)
wb.save(out)
print(out, 'rows', len(rows), 'TOTAL', total)
